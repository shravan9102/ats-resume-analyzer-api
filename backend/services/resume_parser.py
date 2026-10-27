"""Resume file validation and text extraction (PDF, DOCX, TXT).

File type is detected from the file's own bytes (magic numbers) plus the
filename extension, so no native libmagic/DLL is needed (python-magic fails or
misbehaves on many Windows setups).
"""
import io
import zipfile
from typing import Optional, Tuple

import PyPDF2
import pdfplumber
from docx import Document

from backend.core.config import (
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB,
    MAX_PDF_PAGES,
)
from backend.utils.file_utils import (
    FileParsingError,
    FileValidationError,
    TextExtractionError,
    log_error,
    log_info,
    log_warning,
)

__all__ = [
    "FileParsingError",
    "FileValidationError",
    "detect_file_type",
    "validate_file",
    "extract_text",
    "parse_resume_file",
]

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"   # legacy .doc


def detect_file_type(file_data: bytes, filename: str = "") -> Optional[str]:
    """Return 'pdf', 'docx', 'doc', 'txt' or None, based on content."""
    head = file_data[:1024]

    if b"%PDF-" in head:
        return "pdf"

    if file_data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(file_data)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
        return None

    if file_data[:8] == _OLE_MAGIC:
        return "doc"

    # Plain text: must decode cleanly and contain no NUL bytes.
    if filename.lower().endswith(".txt") or not filename:
        if b"\x00" not in file_data[:4096]:
            try:
                file_data[:100000].decode("utf-8")
                return "txt"
            except UnicodeDecodeError:
                try:
                    file_data[:100000].decode("cp1252")
                    return "txt"
                except UnicodeDecodeError:
                    return None
    return None


def validate_file(file_data: bytes, filename: str) -> Tuple[bool, str, Optional[str]]:
    """Return (is_valid, error_message, file_type)."""
    size = len(file_data)

    if size == 0:
        return False, "The uploaded file is empty. Please check the file and try again.", None

    if size > MAX_FILE_SIZE_BYTES:
        return False, (
            f"File size ({size / (1024 * 1024):.2f} MB) exceeds the maximum of "
            f"{MAX_FILE_SIZE_MB} MB. Please upload a smaller file."
        ), None

    file_type = detect_file_type(file_data, filename or "")
    if file_type is None:
        return False, (
            "Unsupported or corrupted file. Please upload a valid PDF, DOCX or TXT resume."
        ), None

    return True, "", file_type


# --------------------------------------------------------------------------- PDF

def _extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))
        for page in list(reader.pages)[:MAX_PDF_PAGES]:
            if "/Annots" not in page:
                continue
            for annot_ref in page["/Annots"]:
                try:
                    annot = annot_ref.get_object()
                    if annot.get("/Subtype") != "/Link":
                        continue
                    uri = annot.get("/A", {}).get("/URI", "")
                    if isinstance(uri, bytes):
                        uri = uri.decode("utf-8", errors="ignore")
                    uri = str(uri).strip()
                    if uri.startswith("http"):
                        urls.append(uri)
                except Exception:
                    continue
    except Exception:
        pass
    return "\n".join(dict.fromkeys(urls))


def _extract_pdf_with_pdfplumber(file_data: bytes) -> str:
    parts = []
    with pdfplumber.open(io.BytesIO(file_data)) as pdf:
        for page in pdf.pages[:MAX_PDF_PAGES]:
            page_text = page.extract_text()
            if page_text:
                parts.append(page_text)
    text = "\n".join(parts)
    if not text.strip():
        raise TextExtractionError("pdfplumber extracted no text")
    return text.strip()


def _extract_pdf_with_pypdf2(file_data: bytes) -> str:
    reader = PyPDF2.PdfReader(io.BytesIO(file_data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise TextExtractionError("PDF is password-protected") from exc
    parts = []
    for page in list(reader.pages)[:MAX_PDF_PAGES]:
        page_text = page.extract_text()
        if page_text:
            parts.append(page_text)
    text = "\n".join(parts)
    if not text.strip():
        raise TextExtractionError("PyPDF2 extracted no text")
    return text.strip()


def extract_text_from_pdf(file_data: bytes) -> str:
    text = None
    errors = []
    for name, func in (("pdfplumber", _extract_pdf_with_pdfplumber),
                       ("PyPDF2", _extract_pdf_with_pypdf2)):
        try:
            text = func(file_data)
            log_info(f"PDF text extracted with {name}", context="resume_parser")
            break
        except Exception as exc:
            errors.append(f"{name}: {exc}")
            log_warning(f"{name} failed: {exc}", context="resume_parser")

    if not text:
        log_error(Exception("; ".join(errors)), context="extract_text_from_pdf")
        raise FileParsingError(
            "Could not extract text from this PDF. It may be corrupted, "
            "password-protected, or a scanned image without selectable text.",
            user_message="No readable text found in the PDF.",
        )

    links = _extract_pdf_hyperlinks(file_data)
    return f"{text}\n{links}".strip() if links else text


# -------------------------------------------------------------------------- DOCX

def extract_text_from_docx(file_data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        text = "\n".join(parts)
    except Exception as exc:
        log_error(exc, context="extract_text_from_docx")
        raise FileParsingError(
            "Could not read this DOCX file. It may be corrupted; try re-saving it or exporting to PDF."
        ) from exc

    if not text.strip():
        raise FileParsingError("No text could be extracted from the document. It may be empty.")

    try:
        for rel in doc.part.rels.values():
            if "hyperlink" in rel.reltype.lower():
                url = getattr(rel, "target_ref", None) or getattr(rel, "_target", "")
                if isinstance(url, str) and url.startswith("http"):
                    text += "\n" + url
    except Exception:
        pass

    return text.strip()


def extract_text_from_txt(file_data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = file_data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        raise FileParsingError("Could not decode the text file.")
    if not text.strip():
        raise FileParsingError("The text file contains no readable text.")
    return text.strip()


def extract_text_from_doc(file_data: bytes) -> str:
    raise FileParsingError(
        "Legacy .doc files are not supported. Please save the resume as .docx or PDF and try again."
    )


def extract_text(file_data: bytes, file_type: str) -> str:
    if file_type == "pdf":
        return extract_text_from_pdf(file_data)
    if file_type == "docx":
        return extract_text_from_docx(file_data)
    if file_type == "txt":
        return extract_text_from_txt(file_data)
    if file_type == "doc":
        return extract_text_from_doc(file_data)
    raise FileValidationError(f"Unsupported file type: {file_type}")


def parse_resume_file(file_data: bytes, filename: str) -> Tuple[str, dict]:
    """Validate the upload and return (text, metadata). Raises FileValidationError / FileParsingError."""
    log_info(f"parsing file: {filename}", context="parse_resume_file")

    is_valid, error_msg, file_type = validate_file(file_data, filename)
    if not is_valid:
        log_warning(f"validation failed for {filename}: {error_msg}", context="parse_resume_file")
        raise FileValidationError(error_msg)

    try:
        text = extract_text(file_data, file_type)
    except (FileParsingError, FileValidationError):
        raise
    except Exception as exc:
        log_error(exc, context="parse_resume_file_extraction")
        raise FileParsingError(
            "An unexpected error occurred while reading the file. Please try another file."
        ) from exc

    log_info(f"Extracted {len(text)} chars from {filename} ({file_type})", context="parse_resume_file")
    return text, {
        "filename": filename,
        "file_type": file_type,
        "file_size_bytes": len(file_data),
        "text_length": len(text),
        "success": True,
    }
